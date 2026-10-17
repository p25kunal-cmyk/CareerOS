"""Tests for ResumeExporter across DOCX, PDF, Markdown, and Text formats."""

import io
import docx
import fitz
import pytest

from src.resume.exporter import ResumeExporter
from src.resume.models import BulletItem, ResumeVersionData, SectionItem


@pytest.fixture
def sample_version():
    """Fixture providing a populated MBA resume version."""
    return ResumeVersionData(
        resume_id="r_test",
        version_number=1,
        version_name="v1.0 (IIM Placement)",
        sections=[
            SectionItem(
                title="Contact Information",
                section_type="contact",
                display_order=1,
                bullets=[
                    BulletItem(content="Ananya Sen", bullet_type="header", display_order=1),
                    BulletItem(content="ananya.sen@iimnagpur.ac.in | +91 9811122233 | Nagpur", bullet_type="contact_item", display_order=2),
                ]
            ),
            SectionItem(
                title="Academic Qualifications",
                section_type="education",
                display_order=2,
                bullets=[
                    BulletItem(content="Degree | Institution | Board | Year | %/CGPA", bullet_type="text"),
                    BulletItem(content="MBA | IIM Nagpur | Autonomous | 2026 | 7.9/10", bullet_type="text"),
                    BulletItem(content="B.E. Computer Science | BITS Pilani | Autonomous | 2022 | 8.3/10", bullet_type="text"),
                ]
            ),
            SectionItem(
                title="Work Experience",
                section_type="experience",
                display_order=3,
                bullets=[
                    BulletItem(content="Deloitte Consulting - Business Technology Analyst", bullet_type="header"),
                    BulletItem(content="Formulated business process re-engineering plan for Fortune 500 retail bank.", bullet_type="bullet", action_verb="formulated"),
                    BulletItem(content="Automated reconciliation dashboard cutting month-end close by 3 days.", bullet_type="bullet", action_verb="automated"),
                ]
            )
        ]
    )


def test_export_docx(sample_version):
    """Test generating a valid DOCX document with tables."""
    docx_bytes = ResumeExporter.export_docx(sample_version, template_id="iim_nagpur")
    assert len(docx_bytes) > 1000

    # Load back with python-docx to verify integrity
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs_text = " ".join(p.text for p in doc.paragraphs)
    assert "Ananya Sen" in paragraphs_text
    assert "WORK EXPERIENCE" in paragraphs_text

    # Verify table was generated for education
    assert len(doc.tables) >= 1
    table_text = " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert "IIM Nagpur" in table_text
    assert "BITS Pilani" in table_text


def test_export_pdf(sample_version):
    """Test generating a valid PDF document."""
    pdf_bytes = ResumeExporter.export_pdf(sample_version, template_id="iim_nagpur")
    assert len(pdf_bytes) > 1000

    # Verify with PyMuPDF
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    assert len(doc) >= 1
    page_text = doc[0].get_text()
    assert "Ananya Sen" in page_text
    assert "WORK EXPERIENCE" in page_text
    doc.close()


def test_export_markdown(sample_version):
    """Test generating Markdown text."""
    md = ResumeExporter.export_markdown(sample_version)
    assert "# Ananya Sen" in md
    assert "## WORK EXPERIENCE" in md
    assert "- Formulated business process" in md


def test_export_text(sample_version):
    """Test generating clean ASCII text."""
    txt = ResumeExporter.export_text(sample_version)
    assert "WORK EXPERIENCE" in txt
    assert "Deloitte Consulting" in txt


def test_export_html(sample_version):
    """Test generating structured, styled HTML resume."""
    html = ResumeExporter.export_html(sample_version, template_id="iim_nagpur", highlight_metrics=True)
    assert "resume-paper" in html
    assert "Ananya Sen" in html
    assert "WORK EXPERIENCE" in html
    assert "Deloitte Consulting" in html
    assert "edu-table" in html  # Education table check


def test_render_pdf_pages_as_images(sample_version):
    """Test rasterizing PDF vector pages into image bytes."""
    images = ResumeExporter.render_pdf_pages_as_images(sample_version, template_id="iim_nagpur", dpi=100)
    assert len(images) >= 1
    assert isinstance(images[0], bytes)
    assert len(images[0]) > 5000  # Non-trivial PNG image


def test_sample_mba_resume_creation(db):
    """Test creating the high-impact MBA sample resume."""
    from src.resume.repository import ResumeRepository
    from src.resume.sample_data import create_sample_mba_resume

    repo = ResumeRepository(db)
    r_id, v_id = create_sample_mba_resume(repo)
    assert r_id is not None
    assert v_id is not None

    resumes = repo.list_resumes()
    assert len(resumes) >= 1
    assert any(r.is_master for r in resumes)

    version_data = repo.get_version(v_id)
    assert version_data is not None
    assert len(version_data.sections) >= 6
