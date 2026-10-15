"""Tests for ResumeParser across text, DOCX, and PDF formats."""

import io
import docx
import fitz
import pytest

from src.resume.parser import ResumeParser


SAMPLE_RESUME_TEXT = """
Rahul Sharma
rahul.sharma@iimnagpur.ac.in | +91 9876543210 | linkedin.com/in/rahulsharma-iim | Nagpur, India

SUMMARY
MBA Candidate at IIM Nagpur with 3 years of software engineering experience. Specializing in Product Management and Strategy.

EDUCATION
MBA | IIM Nagpur | Autonomous | 2026 | 7.85/10
B.Tech Computer Science | NIT Trichy | Autonomous | 2021 | 8.4/10

WORK EXPERIENCE
Senior Software Engineer - Infosys Ltd
- Spearheaded development of cloud microservices reducing latency by 35% for 2M+ active users.
- Led a cross-functional team of 6 engineers to deliver $1.2M enterprise SaaS product on schedule.
- Optimized database indexing and SQL queries, saving $45,000 in monthly AWS infrastructure costs.

SUMMER INTERNSHIPS
Product Management Intern - Amazon
- Formulated Go-To-Market strategy for Alexa Smart Home devices, targeting 15% adoption growth.
- Analyzed customer telemetry data across 50,000 sessions to identify checkout drop-off bottlenecks.

POSITIONS OF RESPONSIBILITY
Coordinator, Product Management Club - IIM Nagpur
- Orchestrated national case competition attracting 1,200+ participants across top B-schools.

SKILLS
Product Strategy, SQL, Python, Wireframing, Agile Scrum, Financial Modeling
"""


def test_parse_text_format():
    """Test parsing standard text/markdown resume."""
    parser = ResumeParser()
    parsed = parser.parse_file(SAMPLE_RESUME_TEXT.encode("utf-8"), "resume.txt")

    assert parsed.contact_info.email == "rahul.sharma@iimnagpur.ac.in"
    assert "9876543210" in parsed.contact_info.phone
    assert "linkedin.com/in/rahulsharma-iim" in parsed.contact_info.linkedin
    assert "Rahul Sharma" in parsed.contact_info.name

    section_titles = [s.title.lower() for s in parsed.sections]
    assert any("education" in t for t in section_titles)
    assert any("experience" in t for t in section_titles)
    assert any("internships" in t for t in section_titles)
    assert any("skills" in t for t in section_titles)

    # Check bullet action verbs and metrics
    exp_sec = next(s for s in parsed.sections if "experience" in s.title.lower())
    assert len(exp_sec.bullets) >= 2

    # Verify action verbs detected
    action_verbs = [b.action_verb for b in exp_sec.bullets if b.action_verb]
    assert "spearheaded" in action_verbs or "led" in action_verbs or "optimized" in action_verbs

    # Verify metrics detected
    all_metrics = [m for b in exp_sec.bullets for m in b.metrics]
    assert any("35%" in m for m in all_metrics)
    assert any("1.2M" in m or "$1.2M" in m for m in all_metrics)


def test_parse_docx_format():
    """Test parsing generated DOCX file with tables and styles in sequence."""
    doc = docx.Document()
    doc.add_heading("Priya Patel", level=0)
    doc.add_paragraph("priya.patel@gmail.com | +91 9123456789 | linkedin.com/in/priyapatel | Mumbai")

    # Add Education
    doc.add_heading("EDUCATION", level=1)
    table = doc.add_table(rows=2, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Degree"
    hdr_cells[1].text = "Institution"
    hdr_cells[2].text = "Year"
    hdr_cells[3].text = "CGPA"

    row_cells = table.rows[1].cells
    row_cells[0].text = "MBA"
    row_cells[1].text = "IIM Nagpur"
    row_cells[2].text = "2026"
    row_cells[3].text = "8.2/10"

    # Add Experience
    doc.add_heading("WORK EXPERIENCE", level=1)
    doc.add_paragraph("Delivered supply chain analytics tool saving ₹25L annually.", style="List Bullet")

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    parser = ResumeParser()
    parsed = parser.parse_file(bio.getvalue(), "priya_resume.docx")

    assert parsed.contact_info.email == "priya.patel@gmail.com"
    assert len(parsed.sections) >= 2

    edu_sec = next(s for s in parsed.sections if "education" in s.title.lower())
    assert any("IIM Nagpur" in b.content for b in edu_sec.bullets)


def test_parse_pdf_format():
    """Test parsing PDF generated with PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)

    page.insert_text(fitz.Point(50, 60), "Amitabh Bachchan", fontsize=16)
    page.insert_text(fitz.Point(50, 80), "amitabh@bollywood.com | +91 9999988888", fontsize=10)
    page.insert_text(fitz.Point(50, 110), "PROFESSIONAL EXPERIENCE", fontsize=12)
    page.insert_text(fitz.Point(50, 130), "• Directed 200+ blockbuster films with over ₹500Cr box office receipts.", fontsize=10)

    pdf_bytes = doc.tobytes()
    doc.close()

    parser = ResumeParser()
    parsed = parser.parse_file(pdf_bytes, "amitabh.pdf")

    assert parsed.contact_info.email == "amitabh@bollywood.com"
    assert len(parsed.sections) >= 1
    exp_sec = next(s for s in parsed.sections if "experience" in s.title.lower())
    assert any("blockbuster" in b.content for b in exp_sec.bullets)
